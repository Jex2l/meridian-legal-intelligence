from pathlib import Path

import pytest
from docx import Document as DocxDocument
from reportlab.lib.pagesizes import LETTER
from reportlab.pdfgen import canvas


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    doc = DocxDocument()
    doc.add_paragraph("MASTER SERVICES AGREEMENT")
    doc.add_paragraph("1. Definitions")
    doc.add_paragraph("As used in this Agreement, the following terms shall have the meanings set forth below.")
    doc.add_paragraph("2. Indemnification")
    doc.add_paragraph("The Seller shall indemnify the Buyer against all losses arising from a breach.")
    doc.add_paragraph("(a) Scope")
    doc.add_paragraph("This indemnity covers direct damages only.")
    path = tmp_path / "sample.docx"
    doc.save(str(path))
    return path


@pytest.fixture
def sample_pdf(tmp_path: Path) -> Path:
    path = tmp_path / "sample.pdf"
    c = canvas.Canvas(str(path), pagesize=LETTER)
    lines = [
        "ARTICLE I",
        "SECTION 1.1 Purpose",
        "This memorandum summarizes the key terms of the proposed transaction.",
        "SECTION 1.2 Background",
        "The parties previously entered into a letter of intent dated January 1, 2024.",
    ]
    y = 700
    for line in lines:
        c.drawString(72, y, line)
        y -= 20
    c.save()
    return path


@pytest.fixture
def blank_scanned_pdf(tmp_path: Path) -> Path:
    """A PDF with no extractable text, simulating a scanned page (used to
    assert the OCR-fallback branch is taken; we don't assert OCR output
    itself here since tesseract may not be installed in CI)."""
    path = tmp_path / "scanned.pdf"
    c = canvas.Canvas(str(path), pagesize=LETTER)
    c.showPage()
    c.save()
    return path
