"""Seed the demo corpus (app/eval/fixtures.py) into Postgres so the eval
harness, and anyone demoing the CLI, has something realistic to query.
Idempotent: skips any document whose filename already exists in the target
workspace/corpus.
"""

import tempfile
from pathlib import Path

from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.db import SessionLocal, init_db
from app.eval.fixtures import CASES, CONTRACTS
from app.ingestion.pipeline import ingest_file
from app.models.models import Document, SourceType, User, Workspace

DEMO_WORKSPACE_NAME = "LexRAG Eval Demo"
DEMO_USER_EMAIL = "demo@lexrag.local"


def _write_docx(sections: list[tuple[str, str]], path: Path) -> None:
    doc = DocxDocument()
    for heading, body in sections:
        doc.add_paragraph(heading)
        doc.add_paragraph(body)
    doc.save(str(path))


def _get_or_create_demo_workspace(session: Session) -> tuple[Workspace, User]:
    ws = session.scalar(select(Workspace).where(Workspace.name == DEMO_WORKSPACE_NAME))
    if ws is None:
        ws = Workspace(name=DEMO_WORKSPACE_NAME)
        session.add(ws)
        session.flush()
    user = session.scalar(select(User).where(User.workspace_id == ws.id, User.email == DEMO_USER_EMAIL))
    if user is None:
        user = User(workspace_id=ws.id, email=DEMO_USER_EMAIL, name="Demo User")
        session.add(user)
        session.commit()
    return ws, user


def _already_ingested(session: Session, filename: str, workspace_id) -> bool:
    return (
        session.scalar(
            select(Document).where(Document.filename == filename, Document.workspace_id == workspace_id)
        )
        is not None
    )


def seed(session: Session, embed: bool = True) -> Workspace:
    ws, user = _get_or_create_demo_workspace(session)

    with tempfile.TemporaryDirectory() as tmp:
        tmp_path = Path(tmp)

        for filename, sections in CONTRACTS.items():
            if _already_ingested(session, filename, ws.id):
                continue
            path = tmp_path / filename
            _write_docx(sections, path)
            ingest_file(
                session, src_path=path, filename=filename,
                workspace_id=ws.id, owner_user_id=user.id, embed=embed,
            )

        for filename, case in CASES.items():
            if _already_ingested(session, filename, None):
                continue
            path = tmp_path / filename
            _write_docx(case["sections"], path)
            ingest_file(
                session, src_path=path, filename=filename,
                workspace_id=None, owner_user_id=None,
                source_type=SourceType.PUBLIC_CORPUS, jurisdiction=case["jurisdiction"], embed=embed,
            )

    return ws


def main() -> None:
    init_db()
    with SessionLocal() as session:
        ws = seed(session)
        print(f"demo workspace_id={ws.id}")


if __name__ == "__main__":
    main()
